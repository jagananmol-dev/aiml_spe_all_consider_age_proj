"""
VEDA AI — Tests for the text, data, and Office parsers

Each test builds a real sample file in a temp directory and runs the
parser on it — no mocks — so the output is what the pipeline will index.
"""

import json
import re
from pathlib import Path

import pytest

from src.workers.ingestion_worker import (
    PARSER_MAP,
    DocumentUploadedEvent,
    parse_csv,
    parse_docx,
    parse_json,
    parse_pptx,
    parse_text,
    read_text_file,
)

REPO_ROOT = Path(__file__).resolve().parents[4]
WEB_FORMATS = REPO_ROOT / "apps" / "web" / "src" / "lib" / "formats.ts"


# ══════════════════════════════════════════════════════
# Format list stays in sync with the web upload page
# ══════════════════════════════════════════════════════


def test_every_web_format_extension_has_a_parser():
    source = WEB_FORMATS.read_text(encoding="utf-8")
    extensions = {
        ext
        for block in re.findall(r"extensions:\s*\[([^\]]*)\]", source)
        for ext in re.findall(r'"([a-z0-9]+)"', block)
    }
    assert extensions, "could not read extensions from formats.ts"
    assert extensions == set(PARSER_MAP), (
        f"web-only: {sorted(extensions - set(PARSER_MAP))}, "
        f"parser-only: {sorted(set(PARSER_MAP) - extensions)}"
    )


# ══════════════════════════════════════════════════════
# Kafka contract with the web app
# ══════════════════════════════════════════════════════


def test_upload_event_accepts_web_camelcase_payload():
    payload = {
        "tenantId": "t1",
        "documentId": "d1",
        "fileName": "report.pdf",
        "fileType": "pdf",
        "storagePath": "documents/report.pdf",
        "storageBucket": "veda-tenant-acme",
        "uploadedBy": "u1",
        "timestamp": "2026-10-08T10:00:00Z",
    }
    event = DocumentUploadedEvent(**payload)
    assert event.tenant_id == "t1"
    assert event.storage_bucket == "veda-tenant-acme"


# ══════════════════════════════════════════════════════
# Plain text and Markdown
# ══════════════════════════════════════════════════════


class TestTextFiles:
    def test_utf8_text(self, tmp_path):
        f = tmp_path / "notes.txt"
        f.write_bytes("Patient: John\r\nDiagnosis: flu".encode("utf-8"))
        result = parse_text(str(f))
        assert result["raw_text"] == "Patient: John\nDiagnosis: flu"
        assert result["word_count"] == 4

    def test_utf8_bom_and_utf16(self, tmp_path):
        bom = tmp_path / "bom.txt"
        bom.write_bytes("ﬁ café".encode("utf-8-sig"))
        utf16 = tmp_path / "u16.txt"
        utf16.write_bytes("Q3 revenue".encode("utf-16"))
        assert read_text_file(str(bom)) == "ﬁ café"
        assert read_text_file(str(utf16)) == "Q3 revenue"

    def test_non_utf8_falls_back(self, tmp_path):
        f = tmp_path / "legacy.txt"
        f.write_bytes("Total: 100€".encode("cp1252"))
        assert parse_text(str(f))["raw_text"] == "Total: 100€"

    def test_markdown_kept_verbatim(self, tmp_path):
        f = tmp_path / "readme.md"
        f.write_text("# Policy\n\nRefunds within 30 days.", encoding="utf-8")
        assert parse_text(str(f))["raw_text"].startswith("# Policy")


# ══════════════════════════════════════════════════════
# CSV
# ══════════════════════════════════════════════════════


class TestCsv:
    def test_rows_become_labelled_records(self, tmp_path):
        f = tmp_path / "accounts.csv"
        f.write_text("account,balance,currency\nACME,1200.50,USD\nGlobex,99,EUR\n", "utf-8")
        result = parse_csv(str(f))
        blocks = result["raw_text"].split("\n\n")
        assert blocks[0] == "account: ACME\nbalance: 1200.50\ncurrency: USD"
        assert len(blocks) == 2
        assert result["tables"][0]["rows"] == 3

    def test_semicolon_delimiter_detected(self, tmp_path):
        f = tmp_path / "eu.csv"
        f.write_text("name;amount\nRent;1.500,00\n", "utf-8")
        assert "amount: 1.500,00" in parse_csv(str(f))["raw_text"]

    def test_empty_csv(self, tmp_path):
        f = tmp_path / "empty.csv"
        f.write_text("", "utf-8")
        assert parse_csv(str(f))["raw_text"] == ""


# ══════════════════════════════════════════════════════
# JSON and JSON Lines
# ══════════════════════════════════════════════════════


class TestJson:
    def test_array_of_records(self, tmp_path):
        f = tmp_path / "patients.json"
        f.write_text(
            json.dumps(
                [
                    {"id": 1, "name": "Asha", "vitals": {"bp": "120/80"}, "tags": ["a", "b"]},
                    {"id": 2, "name": "Ravi", "allergies": [{"drug": "penicillin"}]},
                ]
            ),
            "utf-8",
        )
        result = parse_json(str(f))
        first, second = result["raw_text"].split("\n\n")
        assert first.splitlines() == [
            "Record 1",
            "id: 1",
            "name: Asha",
            "vitals.bp: 120/80",
            "tags: a, b",
        ]
        assert "allergies[0].drug: penicillin" in second
        assert result["page_count"] == 2

    def test_single_object_split_by_top_level_key(self, tmp_path):
        f = tmp_path / "company.json"
        f.write_text(json.dumps({"company": {"name": "Acme"}, "fiscal_year": 2025}), "utf-8")
        result = parse_json(str(f))
        assert result["raw_text"] == "company.name: Acme\n\nfiscal_year: 2025"

    def test_json_lines(self, tmp_path):
        f = tmp_path / "events.jsonl"
        f.write_text('{"event": "login"}\n\n{"event": "logout"}\n', "utf-8")
        result = parse_json(str(f))
        assert result["page_count"] == 2
        assert "event: logout" in result["raw_text"]

    def test_invalid_json_reports_line(self, tmp_path):
        f = tmp_path / "bad.json"
        f.write_text('{"ok": 1}\n{broken\n', "utf-8")
        with pytest.raises(ValueError, match="line 2"):
            parse_json(str(f))


# ══════════════════════════════════════════════════════
# Word and PowerPoint
# ══════════════════════════════════════════════════════


class TestOffice:
    def test_docx_paragraphs_headings_and_tables_in_order(self, tmp_path):
        import docx

        document = docx.Document()
        document.add_heading("Quarterly Report", level=1)
        document.add_paragraph("Revenue grew 12%.")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Region"
        table.cell(0, 1).text = "Revenue"
        table.cell(1, 0).text = "North"
        table.cell(1, 1).text = "5M"
        document.add_paragraph("Outlook is stable.")
        path = tmp_path / "report.docx"
        document.save(path)

        result = parse_docx(str(path))
        parts = result["raw_text"].split("\n\n")
        assert parts == [
            "# Quarterly Report",
            "Revenue grew 12%.",
            "Region | Revenue\nNorth | 5M",
            "Outlook is stable.",
        ]
        assert result["tables"][0]["rows"] == 2

    def test_pptx_slides_and_notes(self, tmp_path):
        from pptx import Presentation

        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = "Treatment Plan"
        slide.placeholders[1].text = "Start therapy in week 2"
        slide.notes_slide.notes_text_frame.text = "Confirm with Dr. Rao"
        deck.slides.add_slide(deck.slide_layouts[5]).shapes.title.text = "Follow-up"
        path = tmp_path / "plan.pptx"
        deck.save(path)

        result = parse_pptx(str(path))
        assert result["page_count"] == 2
        first, second = result["raw_text"].split("\n\n")
        assert first.splitlines()[0] == "Slide 1"
        assert "Treatment Plan" in first
        assert "Notes: Confirm with Dr. Rao" in first
        assert "Follow-up" in second
