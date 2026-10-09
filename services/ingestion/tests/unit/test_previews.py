"""
VEDA AI — Office preview tests (real files built with python-docx, openpyxl, python-pptx)
"""

import docx
import openpyxl
import pytest
from pptx import Presentation

from src.workers import previews


@pytest.fixture
def word_file(tmp_path):
    d = docx.Document()
    d.add_heading("Leave Policy", level=0)
    d.add_heading("Entitlement", level=1)
    d.add_paragraph("Earned leave is 15 days.")
    table = d.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Type", "Days"
    row = table.add_row().cells
    row[0].text, row[1].text = "Casual", "7"
    path = tmp_path / "policy.docx"
    d.save(path)
    return str(path)


def test_docx_keeps_headings_paragraphs_and_tables(word_file):
    blocks = previews.build_preview(word_file, "docx")["blocks"]
    assert blocks[0] == {"type": "heading", "level": 1, "text": "Leave Policy"}
    assert blocks[1] == {"type": "heading", "level": 2, "text": "Entitlement"}
    assert blocks[2] == {"type": "paragraph", "text": "Earned leave is 15 days."}
    assert blocks[3]["type"] == "table" and blocks[3]["rows"] == [["Type", "Days"], ["Casual", "7"]]


def test_xlsx_keeps_every_sheet_and_caps_rows(tmp_path, monkeypatch):
    monkeypatch.setattr(previews, "MAX_ROWS", 3)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Revenue"
    for i in range(5):
        ws.append([f"Centre {i}", i * 10, None])
    wb.create_sheet("Notes").append(["Owner", "Finance"])
    path = tmp_path / "mis.xlsx"
    wb.save(path)

    sheets = previews.build_preview(str(path), "xlsx")["sheets"]
    assert [s["name"] for s in sheets] == ["Revenue", "Notes"]
    assert sheets[0]["rows"] == [["Centre 0", "0"], ["Centre 1", "10"], ["Centre 2", "20"]]
    assert sheets[0]["total_rows"] == 5


def test_pptx_slides_have_title_bullets_and_notes(tmp_path):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Q3 Review"
    body = slide.placeholders[1].text_frame
    body.text = "Revenue up"
    body.add_paragraph().text = "Costs flat"
    slide.notes_slide.notes_text_frame.text = "Speaker note"
    path = tmp_path / "deck.pptx"
    deck.save(path)

    slides = previews.build_preview(str(path), "pptx")["slides"]
    assert slides == [{"number": 1, "title": "Q3 Review", "bullets": ["Revenue up", "Costs flat"],
                       "tables": [], "notes": "Speaker note"}]


def test_other_formats_have_no_stored_preview(tmp_path):
    assert previews.build_preview(str(tmp_path / "notes.txt"), "txt") is None


def test_store_preview_never_raises(word_file, monkeypatch):
    def broken(*args):
        raise RuntimeError("db down")

    monkeypatch.setattr(previews, "save_preview", broken)
    previews.store_preview_safely("db", "t", "d", word_file, "docx")  # must not raise
